import io
from datetime import date

from openpyxl import load_workbook

from inventory_io import export_assets_csv, export_assets_xlsx, export_table_xlsx, import_assets
from inventory_io.tests.helpers import asset, csv_bytes

TODAY = date(2026, 10, 1)


def sample_rows():
    result = import_assets(csv_bytes([
        asset(),
        asset(**{
            "Asset ID": "0002", "Status": "On_Loan", "Borrower ID": "3712345",
            "Checkout Date": "2026-09-20", "Expected Return Date": "2026-10-04",
        }),
    ]), "x.csv", today=TODAY)
    assert result.ok
    return result.rows


def test_xlsx_export_round_trips_through_the_importer():
    rows = sample_rows()
    again = import_assets(export_assets_xlsx(rows), "export.xlsx", today=TODAY)
    assert again.ok, [str(i) for i in again.issues]
    assert [r.as_dict() for r in again.rows] == [r.as_dict() for r in rows]


def test_csv_export_round_trips_through_the_importer():
    rows = sample_rows()
    data = export_assets_csv(rows).encode("utf-8-sig")
    again = import_assets(data, "export.csv", today=TODAY)
    assert [r.as_dict() for r in again.rows] == [r.as_dict() for r in rows]


def test_xlsx_dates_are_real_dates():
    sheet = load_workbook(io.BytesIO(export_assets_xlsx(sample_rows()))).active
    assert sheet["E2"].value.date() == date(2022, 9, 1)
    assert sheet["E2"].number_format == "yyyy-mm-dd"


def test_xlsx_never_writes_formulas():
    data = export_table_xlsx(["Name"], [["=HYPERLINK(\"http://evil\",\"click\")"]])
    sheet = load_workbook(io.BytesIO(data)).active
    assert sheet["A2"].data_type == "s"
    assert sheet["A2"].value.startswith("=HYPERLINK")


def test_csv_escapes_formula_like_text():
    row = {
        "asset_id": "0001", "name": "=cmd|' /C calc'!A0", "category": "Kits",
        "type": "Basic Kit", "acquisition_date": date(2022, 9, 1), "status": "Available",
    }
    text = export_assets_csv([row])
    assert ",'=cmd" in text
    again = import_assets(text.encode(), "x.csv", today=TODAY)
    assert again.rows[0].name == "=cmd|' /C calc'!A0"


def test_export_accepts_plain_dicts():
    text = export_assets_csv([{"asset_id": "0009", "name": "Kit"}])
    assert text.splitlines()[1].startswith("0009,Kit,")
