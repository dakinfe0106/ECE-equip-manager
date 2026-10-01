"""Writing rows back out as .xlsx or .csv.

The asset export uses the same headers the importer accepts, so a file can be
exported, edited in Excel and imported again.
"""

import csv
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from . import spec
from .normalize import FORMULA_PREFIXES

DATE_FORMAT = "yyyy-mm-dd"


def escape_csv_cell(value):
    """Stop text like '=HYPERLINK(...)' running as a formula when the CSV is opened.

    Excel shows the leading apostrophe; our importer strips it again.
    """
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def export_table_xlsx(headers, rows, sheet_title="Export") -> bytes:
    """Write any table (assets, analytics, borrowers) to .xlsx bytes."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title[:31]  # Excel's limit
    sheet.append(list(headers))
    for values in rows:
        sheet.append([None for _ in headers])
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=sheet.max_row, column=column)
            cell.value = value
            if isinstance(value, str):
                # openpyxl treats strings starting with "=" as formulas; force plain text.
                cell.data_type = "s"
            elif isinstance(value, (date, datetime)):
                cell.number_format = DATE_FORMAT
    sheet.freeze_panes = "A2"
    for column, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = max(12, len(str(header)) + 2)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_table_csv(headers, rows) -> str:
    """Write any table to CSV text. Encode as 'utf-8-sig' so Excel reads accents."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow([escape_csv_cell(h) for h in headers])
    for values in rows:
        writer.writerow([_csv_value(v) for v in values])
    return buffer.getvalue()


def _csv_value(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    return escape_csv_cell(value)


def _asset_values(assets):
    for asset in assets:
        data = asset if isinstance(asset, dict) else asset.as_dict()
        yield [data.get(fld) for fld in spec.FIELDS]


def export_assets_xlsx(assets) -> bytes:
    """Export AssetRow objects (or dicts with the same keys) to .xlsx bytes."""
    return export_table_xlsx(list(spec.EXPORT_HEADERS.values()), _asset_values(assets), "Inventory")


def export_assets_csv(assets) -> str:
    """Export AssetRow objects (or dicts with the same keys) to CSV text."""
    return export_table_csv(list(spec.EXPORT_HEADERS.values()), _asset_values(assets))
