"""Reading uploaded .csv and .xlsx files into raw rows.

Nothing here knows what the columns mean; it only produces
(row_number, [cell values]) pairs plus the header row.
"""

import csv
import io
from dataclasses import dataclass, field

from openpyxl import load_workbook

from . import spec
from .issues import Issue

_XLSX_MAGIC = b"PK\x03\x04"


class FileRejected(Exception):
    """The file as a whole cannot be read; ``issue`` explains why."""

    def __init__(self, message: str):
        super().__init__(message)
        self.issue = Issue(message)


@dataclass
class RawSheet:
    headers: list
    rows: list = field(default_factory=list)  # (row_number, [values])


def read_upload(data: bytes, filename: str, sheet_name: str | None = None) -> RawSheet:
    """Read uploaded bytes. Raises FileRejected if the file cannot be used."""
    if not data:
        raise FileRejected("The file is empty.")
    if len(data) > spec.MAX_FILE_BYTES:
        limit_mb = spec.MAX_FILE_BYTES // (1024 * 1024)
        raise FileRejected(f"The file is larger than the {limit_mb} MB limit.")

    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension == "xlsx":
        if not data.startswith(_XLSX_MAGIC):
            raise FileRejected("The file is named .xlsx but is not an Excel workbook.")
        sheet = _read_xlsx(data, sheet_name)
    elif extension == "csv":
        if data.startswith(_XLSX_MAGIC):
            raise FileRejected("The file is named .csv but is an Excel workbook; rename it to .xlsx.")
        sheet = _read_csv(data)
    elif extension == "xls":
        raise FileRejected("Old .xls files are not supported; save the file as .xlsx in Excel first.")
    else:
        raise FileRejected("Only .xlsx and .csv files can be imported.")

    if len(sheet.rows) > spec.MAX_ROWS:
        raise FileRejected(f"The file has more than {spec.MAX_ROWS} rows.")
    return sheet


def _read_xlsx(data: bytes, sheet_name: str | None) -> RawSheet:
    try:
        # read_only streams the file; data_only returns formula results, not formulas.
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many different exception types
        raise FileRejected("The Excel file could not be opened; it may be damaged.") from exc
    try:
        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise FileRejected(f"The workbook has no sheet named '{sheet_name}'.")
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.worksheets[0]
        return _to_sheet(enumerate(worksheet.iter_rows(values_only=True), start=1))
    finally:
        workbook.close()


def _read_csv(data: bytes) -> RawSheet:
    try:
        text = data.decode("utf-8-sig")  # strips the BOM Excel adds to "CSV UTF-8"
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")  # Excel's default "CSV" on Windows
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return _to_sheet(enumerate(csv.reader(io.StringIO(text), dialect), start=1))


def _to_sheet(numbered_rows) -> RawSheet:
    """Use the first non-empty row as the header; skip fully blank rows."""
    sheet = None
    for number, values in numbered_rows:
        values = list(values)
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        if sheet is None:
            sheet = RawSheet(headers=values)
        else:
            sheet.rows.append((number, values))
    if sheet is None:
        raise FileRejected("The file has no header row.")
    return sheet
